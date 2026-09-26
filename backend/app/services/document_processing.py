import io
import re
from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from pymongo.database import Database

from app.models import DocumentChunkDocument, StudyMaterialDocument


class DocumentProcessingError(ValueError):
    pass


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def extract_pages(filename: str, content: bytes) -> list[tuple[int, str]]:
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix == "txt":
        return [(1, clean_text(content.decode("utf-8")))]
    if suffix == "pdf":
        return [(index + 1, clean_text(page.extract_text() or "")) for index, page in enumerate(PdfReader(io.BytesIO(content)).pages)]
    if suffix == "docx":
        return [(1, clean_text(" ".join(item.text for item in Document(io.BytesIO(content)).paragraphs)))]
    if suffix == "pptx":
        return [(index + 1, clean_text(" ".join(shape.text for shape in slide.shapes if hasattr(shape, "text")))) for index, slide in enumerate(Presentation(io.BytesIO(content)).slides)]
    raise DocumentProcessingError("Unsupported document type.")


def chunk_pages(pages: list[tuple[int, str]], chunk_size: int = 1000, overlap: int = 150) -> list[dict]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("chunk_size must be positive and overlap must be smaller than chunk_size.")
    chunks: list[dict] = []
    for page, text in pages:
        for start in range(0, len(text), chunk_size - overlap):
            value = text[start:start + chunk_size]
            if value:
                chunks.append({"page_number": page, "content": value})
            if start + chunk_size >= len(text):
                break
    return chunks


class DocumentProcessingService:
    allowed_extensions = {"pdf", "docx", "pptx", "txt"}
    allowed_content_types = {
        "pdf": {"application/pdf"},
        "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        "pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation"},
        "txt": {"text/plain", "application/octet-stream"},
    }

    def __init__(self, database: Database, storage_path: Path, max_upload_bytes: int) -> None:
        self.database = database
        self.storage_path = storage_path
        self.max_upload_bytes = max_upload_bytes

    def upload(
        self,
        *,
        subject_id: str | None,
        filename: str,
        content_type: str,
        content: bytes,
        course_id: str | None = None,
        syllabus_id: str | None = None,
        topic_id: str | None = None,
    ) -> StudyMaterialDocument:
        extension = self._validate_upload(filename, content_type, content)
        resolved_subject_id: str | None = None
        if subject_id:
            subject = self.database.subjects.find_one({"_id": subject_id})
            if subject is None:
                subject = self.database.subjects.find_one({"code": subject_id.upper()})
            if subject is None:
                raise DocumentProcessingError("Subject not found.")
            resolved_subject_id = str(subject["_id"])

        storage_key = f"{uuid4()}.{extension}"
        self.storage_path.mkdir(parents=True, exist_ok=True)
        (self.storage_path / storage_key).write_bytes(content)
        material = StudyMaterialDocument(
            subject_id=resolved_subject_id,
            course_id=course_id,
            syllabus_id=syllabus_id,
            topic_id=topic_id,
            filename=Path(filename).name,
            storage_key=storage_key,
            content_type=content_type,
            size_bytes=len(content),
        )
        self.database.study_materials.insert_one(material.model_dump(by_alias=True))
        return material

    def process(self, material_id: str) -> tuple[StudyMaterialDocument, int]:
        stored = self.database.study_materials.find_one({"_id": material_id})
        if stored is None:
            stored = self.database.study_materials.find_one({"id": material_id})
        if not stored:
            raise DocumentProcessingError("Material not found.")
        material = StudyMaterialDocument.model_validate(stored)
        path = self.storage_path / material.storage_key
        if not path.is_file():
            raise DocumentProcessingError("Stored material is unavailable.")

        try:
            pages = extract_pages(material.filename, path.read_bytes())
            chunks = chunk_pages(pages)
        except Exception as error:
            self.database.study_materials.update_one({"_id": material.id}, {"$set": {"status": "failed"}})
            raise DocumentProcessingError("The material could not be processed.") from error

        self.database.document_chunks.delete_many({"study_material_id": material.id})
        for index, chunk in enumerate(chunks):
            document = DocumentChunkDocument(
                study_material_id=material.id,
                chunk_index=index,
                page_number=chunk["page_number"],
                content=chunk["content"],
                metadata={
                    "source_file": material.filename,
                    "study_material_id": material.id,
                    **({"subject_id": material.subject_id} if material.subject_id else {}),
                    **({"course_id": material.course_id} if material.course_id else {}),
                    **({"syllabus_id": material.syllabus_id} if material.syllabus_id else {}),
                    **({"topic_id": material.topic_id} if material.topic_id else {}),
                },
            )
            self.database.document_chunks.insert_one(document.model_dump(by_alias=True))
        self.database.study_materials.update_one({"_id": material.id}, {"$set": {"status": "processed"}})
        material.status = "processed"
        return material, len(chunks)

    def get_status(self, material_id: str) -> dict:
        material = self.database.study_materials.find_one({"_id": material_id})
        if material is None:
            material = self.database.study_materials.find_one({"id": material_id})
        if not material:
            raise DocumentProcessingError("Material not found.")
        resolved_material_id = str(material.get("id", material.get("_id", material_id)))
        indexed_chunks = list(
            self.database.document_chunks.find(
                {
                    "study_material_id": resolved_material_id,
                    "embedding": {"$exists": True, "$ne": None},
                },
                {"embedding_model": 1},
            )
        )
        embedding_models = {item.get("embedding_model") for item in indexed_chunks if item.get("embedding_model")}
        return {
            "material": StudyMaterialDocument.model_validate(material),
            "chunk_count": self.database.document_chunks.count_documents(
                {"study_material_id": resolved_material_id}
            ),
            "indexed_chunk_count": len(indexed_chunks),
            "embedding_model": next(iter(embedding_models), None) if len(embedding_models) == 1 else None,
        }

    def _validate_upload(self, filename: str, content_type: str, content: bytes) -> str:
        extension = Path(filename).suffix.lower().removeprefix(".")
        if extension not in self.allowed_extensions:
            raise DocumentProcessingError("Only PDF, DOCX, PPTX and TXT files are supported.")
        if content_type not in self.allowed_content_types[extension]:
            raise DocumentProcessingError("The file content type does not match its extension.")
        if not content:
            raise DocumentProcessingError("The uploaded file is empty.")
        if len(content) > self.max_upload_bytes:
            raise DocumentProcessingError("The uploaded file is too large.")
        return extension
