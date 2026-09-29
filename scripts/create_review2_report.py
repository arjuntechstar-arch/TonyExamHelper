from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Review_2_Technical_Implementation_Report.docx")
NAVY = "102A43"
BLUE = "1768B4"
CYAN = "2FC2D2"
GRAY = "52606D"
PALE = "EAF3F8"


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def cell_text(cell, value: str, bold: bool = False, color: str = "000000") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(value)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, NAVY)
        cell_text(cell, header, True, "FFFFFF")
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 0:
                shade(cells[index], PALE)
            cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph()


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def heading(doc, value: str, level: int = 1) -> None:
    p = doc.add_heading(value, level=level)
    for run in p.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("AI QUESTION GENERATION SYSTEM")
r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(24); r.font.color.rgb = RGBColor.from_string(NAVY)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Review 2 — Technical Implementation and Intermediate Results Report")
r.font.name = "Aptos"; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(BLUE)
doc.add_paragraph()

summary = doc.add_table(rows=2, cols=3)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (label, value) in enumerate((("Review milestone", "Review 2"), ("Implementation status", "50% complete"), ("System focus", "RAG-based assessment generation"))):
    shade(summary.rows[0].cells[i], NAVY); cell_text(summary.rows[0].cells[i], label, True, "FFFFFF")
    shade(summary.rows[1].cells[i], PALE); cell_text(summary.rows[1].cells[i], value, True, NAVY)
doc.add_paragraph()

heading(doc, "1. Executive Summary")
doc.add_paragraph(
    "The project implements a Retrieval-Augmented Generation (RAG) pipeline that transforms uploaded study material into grounded assessment questions. "
    "The current system combines document preprocessing, vector retrieval, hybrid re-ranking, LLM-based question construction, deterministic quality checks, "
    "optional LLM-based verification, provider failover, and request-level monitoring."
)
doc.add_paragraph(
    "The main Review 2 outcome is an operational end-to-end pipeline: a user can upload study material, select a blueprint, generate a complete question paper, "
    "and inspect every generation stage and failure reason through the Monitoring page."
)

heading(doc, "2. Problem Statement and Objectives")
add_bullets(doc, [
    "Generate assessment questions from source material rather than relying on generic, ungrounded LLM output.",
    "Create varied questions from the same material by rotating factual coverage and assessment angles.",
    "Validate structure, answerability, source faithfulness, duplicate risk, and distractor quality before publishing a question.",
    "Provide transparent handling of model unavailability, provider rate limits, invalid model output, and validation failures.",
])

heading(doc, "3. Implemented Architecture")
doc.add_paragraph("The application uses a custom Python orchestration graph. It is not built on LangChain or LangGraph; agent-like stages are implemented directly in the backend.")
add_table(doc, ["Phase", "Implemented component / technique", "Purpose"], [
    ["1. Ingestion", "PDF, DOCX, PPTX, and TXT extraction", "Converts uploaded material into page-level text."],
    ["2. Cleaning", "Regex noise filtering and repeated-header/footer removal", "Removes emails, copyright/licensing text, opaque identifiers, and repeated extraction noise."],
    ["3. Chunking", "Fixed-size sliding window: 1,000 characters, 150-character overlap, page bounded", "Preserves local context while producing retrievable units. This is not semantic chunking yet."],
    ["4. Indexing", "hash-embedding-v1 (256 dimensions) stored with MongoDB chunk records", "Produces deterministic development vectors for each chunk."],
    ["5. Retrieval", "Cosine similarity over stored vectors", "Selects an initial source-evidence candidate pool."],
    ["6. Re-ranking", "Hybrid score: 75% vector score + 25% lexical overlap; MMR diversification", "Improves keyword coverage and reduces redundant retrieved chunks."],
    ["7. Evidence extraction", "Complete-sentence fact extraction and source-noise filters", "Supplies bounded, usable fact records to generation prompts."],
    ["8. Generation", "Hosted LLM, constrained JSON schema, pattern/difficulty/Bloom prompts", "Constructs source-grounded question candidates."],
    ["9. Validation", "Rule-based schema, pattern, quality, duplicate, and source-faithfulness checks", "Rejects malformed, generic, duplicate, or unsupported questions."],
    ["10. Observability", "Generation run IDs, persisted stage events, Monitoring UI", "Makes each request and failure reason auditable."],
], [1.05, 3.4, 2.75])

heading(doc, "4. Models and Provider Strategy")
add_table(doc, ["Area", "Current implementation"], [
    ["Generation provider order", "OpenRouter route 1 → OpenRouter route 2 → OpenRouter route 3 → NVIDIA NIM fallback."],
    ["Configured NVIDIA model", "moonshotai/kimi-k3, used after all configured OpenRouter routes are unavailable or rate-limited."],
    ["Rate-limit policy", "HTTP 429 places a route in a shared cooldown. Retry-After is respected when provided; otherwise the default cooldown is 60 seconds (maximum 5 minutes)."],
    ["Model health", "The navigation bar displays API status and model-route readiness."],
    ["Independent solver and critic", "Implemented as separate LLM calls. A genuinely independent second model/provider should be configured through LLM_CRITIC_MODEL for production evaluation independence."],
    ["Local deterministic provider", "Available for controlled/offline testing only; not presented as a production-quality LLM critic."],
], [1.9, 5.3])

heading(doc, "5. Question Paper Generation Workflow")
doc.add_paragraph("A paper request honours the complete selected template blueprint. For example, CT 2 is configured with 10 MCQs and 5 short-answer items; therefore one request generates one complete paper across both sections.")
add_table(doc, ["Trace event", "Meaning"], [
    ["Blueprint section 1: 10 MCQ questions", "The first section requires ten question slots."],
    ["Preparing candidate 1/10", "Question slot 1 of the first section is being prepared; it is not an answer option or a separate template."],
    ["Model attempt 1", "First generation attempt for that question slot."],
    ["Pattern / question validation", "Checks template conformity, MCQ structure, duplicate risk, quality, and source grounding."],
    ["LLM critic / independent solver", "Optional evaluation calls inspect faithfulness and solve the MCQ without receiving the answer key."],
    ["Model attempt 2", "A retry for the same question slot only if generation or validation fails."],
], [2.6, 4.6])

heading(doc, "6. Intermediate Results Obtained")
add_bullets(doc, [
    "Complete material-to-question-paper workflow is operational: upload, parse, clean, chunk, index, retrieve, generate, validate, and display results.",
    "Hybrid retrieval and MMR diversification are implemented to improve contextual coverage and reduce redundant evidence.",
    "Question candidates are validated individually, enabling targeted retries rather than restarting an entire paper for one failed item.",
    "Provider failover automatically advances from rate-limited OpenRouter routes to the next route, with NVIDIA NIM as the final fallback.",
    "Critic quality scores submitted on a 0–10 scale are normalized to the required 0–1 scale; invalid critic/solver responses are retried before regenerating a question.",
    "Monitoring captures queued, retrieval, blueprint, preparation, model, failover, validation, critic, solver, retry, completion, and failure events.",
])

heading(doc, "7. Validation and Monitoring")
add_table(doc, ["Validation layer", "Checks performed"], [
    ["Structural validation", "Required fields, question type, template pattern, marks, MCQ option count, and correct-answer reference."],
    ["Quality validation", "Duplicate questions/options, generic distractors, malformed stems, explanation quality, and marks-to-complexity alignment."],
    ["Source validation", "Source availability, source noise, factual grounding, and excessive source copying."],
    ["LLM critic", "Groundedness, answerability, single correct answer, completeness, distractor plausibility, Bloom/difficulty/mark alignment, and quality score."],
    ["Independent solver", "Answers an MCQ using evidence without seeing the generated answer key; disagreement rejects the candidate."],
], [2.05, 5.15])

heading(doc, "8. Completion Status — 50%")
add_table(doc, ["Completed", "Planned / in progress"], [
    ["Document ingestion, cleaning, and fixed-size overlapped chunking", "Sentence/semantic chunking experiments and parameter tuning"],
    ["Hash embeddings, MongoDB storage, cosine retrieval, lexical scoring, and MMR", "Production semantic embeddings, HNSW/vector index, and retrieval re-ranking evaluation"],
    ["Hosted LLM generation, provider failover, candidate retries, and validation", "Domain-specific fine-tuning and advanced question-pattern tuning"],
    ["Generation trace persistence, Monitoring page, API/model status", "RAG metrics dashboard, semantic cache, long-term memory policy, and load testing"],
    ["Critic and solver integration", "Dedicated second provider/model configuration for independent production evaluation"],
], [3.6, 3.6])

heading(doc, "9. Expected Outcomes")
add_bullets(doc, [
    "Generate multiple distinct, source-grounded questions from the same study material through evidence rotation, question-pattern rotation, and retry diversity.",
    "Generate a complete paper for the selected blueprint, retaining its section counts, marks, question types, and patterns.",
    "Reduce failed and slow requests by skipping rate-limited model routes during their cooldown period.",
    "Provide administrators with exact request-level diagnostics for retrieval, model, validation, rate-limit, and provider failures.",
    "Measure future RAG quality using context relevance, faithfulness, retrieval precision/recall, duplicate rate, validation pass rate, latency, and provider error rate." ,
])

heading(doc, "10. Conclusion")
doc.add_paragraph(
    "The Review 2 implementation demonstrates an end-to-end RAG assessment engine with grounded retrieval, structured generation, layered validation, provider resilience, and operational traceability. "
    "The next milestone focuses on replacing baseline embeddings and exhaustive vector scanning with production retrieval infrastructure, completing independent multi-model evaluation, and exposing measurable RAG quality metrics in the UI."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("AI Question Generation System • Review 2 Technical Report")
run.font.name = "Aptos"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(GRAY)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
