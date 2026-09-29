from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("docs/AI_Question_Generation_Review_Presentation.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(12, 29, 56)
BLUE = RGBColor(25, 104, 180)
CYAN = RGBColor(47, 194, 210)
ORANGE = RGBColor(244, 154, 57)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(31, 47, 64)
MUTED = RGBColor(91, 110, 128)
PALE = RGBColor(240, 246, 251)
GREEN = RGBColor(37, 150, 99)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True; frame.margin_left = 0; frame.margin_right = 0
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = value; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=16, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = Inches(0.05); frame.margin_right = Inches(0.05)
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = item; p.level = 0; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(10); p.bullet = True
    return box


def base(title, subtitle=None, page=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 0.22, CYAN)
    text(slide, title, 0.62, 0.45, 11.8, 0.48, 27, NAVY, True)
    if subtitle:
        text(slide, subtitle, 0.64, 0.98, 11.8, 0.35, 11.5, MUTED)
    if page:
        text(slide, str(page), 12.3, 7.03, 0.45, 0.2, 10, MUTED, False, PP_ALIGN.RIGHT)
    return slide


def pill(slide, label, x, y, w, color=BLUE):
    rect(slide, x, y, w, 0.32, color, True)
    text(slide, label, x + 0.08, y + 0.06, w - 0.16, 0.18, 10, WHITE, True, PP_ALIGN.CENTER)


# 1. Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
rect(slide, 0, 0, 13.333, 7.5, NAVY)
rect(slide, 0, 0, 0.24, 7.5, CYAN)
text(slide, "AI Question Generation System", 0.86, 1.45, 10.9, 0.7, 34, WHITE, True)
text(slide, "Technical Implementation Review", 0.88, 2.24, 8.8, 0.42, 20, CYAN, True)
text(slide, "RAG-based assessment generation, validation, observability, and quality assurance", 0.88, 2.93, 10.9, 0.48, 17, RGBColor(211, 225, 239))
rect(slide, 0.88, 4.2, 3.0, 0.72, BLUE, True)
text(slide, "Progress: 50%", 1.08, 4.41, 2.6, 0.25, 18, WHITE, True, PP_ALIGN.CENTER)
text(slide, "Project Review • Tomorrow", 0.9, 6.65, 5, 0.25, 13, RGBColor(211, 225, 239))

# 2. objective
slide = base("Project Objective", "Automate creation of reliable and diverse examination questions from study material.", 2)
rect(slide, 0.64, 1.62, 12.0, 4.73, PALE, True)
text(slide, "Core problem", 1.03, 2.03, 2.1, 0.3, 19, NAVY, True)
text(slide, "Generating many meaningful questions from the same material without repetition, hallucination, or weak answer options.", 1.03, 2.46, 10.7, 0.63, 21, INK)
text(slide, "System goals", 1.03, 3.56, 2.1, 0.3, 19, NAVY, True)
bullets(slide, ["Source-grounded question generation using RAG.", "Multiple cognitive patterns: conceptual, application, scenario-based, and numerical.", "Question-by-question validation, duplicate prevention, and transparent error reporting.", "End-to-end monitoring for retrieval, generation, validation, and provider failures."], 1.04, 3.99, 10.8, 1.9, 17)

# 3 architecture
slide = base("End-to-End Technical Architecture", "Custom orchestration workflow; the application does not depend on LangChain or LangGraph.", 3)
steps = [
    ("Upload", "PDF / study material"), ("Preprocess", "clean + chunk"), ("Index", "embeddings + HNSW"),
    ("Retrieve", "hybrid + MMR"), ("Extract", "facts / evidence"), ("Generate", "LLM MCQs"),
    ("Validate", "rules + critic"), ("Monitor", "traces + UI")]
for i, (heading, detail) in enumerate(steps):
    x = 0.52 + (i % 4) * 3.18; y = 1.66 + (i // 4) * 2.28
    rect(slide, x, y, 2.56, 1.27, WHITE, True, RGBColor(207, 221, 233))
    text(slide, f"{i+1:02d}", x+0.18, y+0.18, 0.4, 0.25, 13, CYAN, True)
    text(slide, heading, x+0.18, y+0.52, 2.1, 0.28, 18, NAVY, True)
    text(slide, detail, x+0.18, y+0.86, 2.13, 0.22, 12, MUTED)
    if i % 4 != 3:
        text(slide, "→", x+2.68, y+0.51, 0.32, 0.28, 20, BLUE, True, PP_ALIGN.CENTER)
text(slide, "Each request is traced across every phase and returns either validated questions, partial results, or an exact failure reason.", 0.72, 6.43, 11.9, 0.34, 16, GREEN, True, PP_ALIGN.CENTER)

# 4 models
slide = base("Models and Providers", "Provider-agnostic design: generation model is chosen through environment configuration.", 4)
rect(slide, 0.65, 1.48, 5.85, 4.9, PALE, True)
text(slide, "Current deployment", 0.98, 1.86, 4.8, 0.32, 20, NAVY, True)
pill(slide, "ACTIVE PROVIDER", 0.98, 2.35, 1.48, CYAN)
text(slide, "NVIDIA NIM", 0.98, 2.83, 4.55, 0.43, 26, BLUE, True)
text(slide, "Configured generation model", 0.98, 3.39, 4.5, 0.25, 13, MUTED)
text(slide, "moonshotai/kimi-k3", 0.98, 3.72, 4.6, 0.39, 22, INK, True)
text(slide, "The UI health indicator displays API connectivity and model readiness.", 0.98, 4.53, 4.68, 0.54, 15, INK)
rect(slide, 6.82, 1.48, 5.85, 4.9, WHITE, True, RGBColor(207, 221, 233))
text(slide, "Supported provider paths", 7.16, 1.86, 4.8, 0.32, 20, NAVY, True)
bullets(slide, ["NVIDIA NIM — current hosted model path.", "OpenRouter — configurable hosted model path.", "OpenAI-compatible endpoint — configurable model path.", "Deterministic provider — controlled testing only; not an LLM critic or production-quality generator.", "Independent solver and LLM critic — invoked when a separately configured capable model is available."], 7.14, 2.38, 4.85, 3.25, 15.5)

# 5 algorithms by phase
slide = base("Algorithms and Techniques by Phase", "The pipeline combines information retrieval, constrained LLM generation, and deterministic quality controls.", 5)
rows = [
    ("Preprocessing", "Text normalization, regex noise filtering, chunking with overlap", "Creates clean, semantically coherent source units"),
    ("Indexing", "Dense embeddings + HNSW approximate nearest-neighbor index", "Enables scalable semantic retrieval"),
    ("Retrieval", "Hybrid dense/lexical retrieval, cosine similarity, MMR", "Finds relevant evidence while reducing redundancy"),
    ("Generation", "Prompt engineering, constrained JSON output, diversity rotation", "Creates different MCQ styles from selected evidence"),
    ("Validation", "Schema/rule checks, regex heuristics, semantic duplicate checks", "Rejects malformed, generic, or repeated questions"),
    ("Evaluation", "Independent blind solver + LLM-as-a-Judge when configured", "Checks answer correctness, ambiguity, and faithfulness")]
for i, row in enumerate(rows):
    y = 1.58 + i*0.77
    rect(slide, 0.65, y, 11.95, 0.59, PALE if i%2==0 else WHITE, False)
    text(slide, row[0], 0.85, y+0.15, 1.72, 0.2, 13.5, NAVY, True)
    text(slide, row[1], 2.72, y+0.12, 4.25, 0.28, 12.5, INK)
    text(slide, row[2], 7.22, y+0.12, 5.05, 0.3, 12.5, MUTED)
text(slide, "Phase", 0.85, 1.25, 1.4, 0.2, 12, BLUE, True)
text(slide, "Algorithm / Technique", 2.72, 1.25, 3.2, 0.2, 12, BLUE, True)
text(slide, "Technical purpose", 7.22, 1.25, 3.2, 0.2, 12, BLUE, True)

# 6 generation quality
slide = base("Question Generation and Quality-Control Workflow", "Generation is evidence-led and validated per candidate rather than only after the complete batch.", 6)
blocks = [
    ("Evidence extraction", "Select complete, usable facts from retrieved chunks."),
    ("Diversity planning", "Rotate definition, comparison, application, scenario, and formula/problem patterns."),
    ("LLM generation", "Generate structured four-option MCQs from fact records."),
    ("Candidate validation", "Check structure, answer key, explanation, complexity, source grounding, and distractors."),
    ("Retry or accept", "Retry only failed/duplicate candidate with a changed diversity angle.")]
for i, (h, d) in enumerate(blocks):
    x = 0.55 + i*2.54
    rect(slide, x, 2.0, 2.13, 2.42, WHITE, True, RGBColor(207,221,233))
    rect(slide, x, 2.0, 2.13, 0.14, ORANGE)
    text(slide, str(i+1), x+0.21, 2.29, 0.32, 0.28, 18, ORANGE, True)
    text(slide, h, x+0.2, 2.77, 1.7, 0.56, 15, NAVY, True)
    text(slide, d, x+0.2, 3.49, 1.74, 0.57, 12, INK)
    if i < 4: text(slide, "→", x+2.17, 3.00, 0.3, 0.3, 19, BLUE, True, PP_ALIGN.CENTER)
text(slide, "Quality gates: exactly 4 options • one valid answer • no duplicate options • no generic distractors • source faithfulness • no semantic duplicate", 0.86, 5.22, 11.6, 0.55, 16, GREEN, True, PP_ALIGN.CENTER)

# 7 intermediate
slide = base("Intermediate Results Obtained", "Implemented capabilities and operational improvements achieved so far.", 7)
left = ["Full RAG pipeline: material ingestion → retrieval → generated assessments.", "Hybrid retrieval improves terminology coverage beyond vector search alone.", "Candidate-level retry reduces duplicate failures without restarting the whole request.", "Partial-result aggregation returns valid questions when only a few candidates fail."]
right = ["Structured failures distinguish model unavailable, provider rate limit, retrieval error, generation error, and validation rejection.", "End-to-end run traces capture timestamps, stages, error messages, and final output state.", "Monitoring UI supports request-level diagnosis.", "Navigation displays API connection and configured model readiness."]
rect(slide, 0.66, 1.58, 5.82, 4.9, PALE, True)
rect(slide, 6.85, 1.58, 5.82, 4.9, WHITE, True, RGBColor(207,221,233))
text(slide, "Generation & Retrieval", 1.02, 1.96, 4.4, 0.31, 19, NAVY, True)
bullets(slide, left, 1.0, 2.45, 4.93, 3.3, 15.5)
text(slide, "Reliability & Observability", 7.21, 1.96, 4.7, 0.31, 19, NAVY, True)
bullets(slide, right, 7.2, 2.45, 4.93, 3.3, 15.5)

# 8 monitoring
slide = base("Observability and Failure Diagnostics", "Every generation request is represented by a run ID and stage-level event trace.", 8)
trace = [("Request started", CYAN), ("Retrieval", BLUE), ("Evidence extraction", BLUE), ("Generation", ORANGE), ("Validation", ORANGE), ("Completed / Failed", GREEN)]
for i, (name, color) in enumerate(trace):
    x = 0.62 + i*2.08
    rect(slide, x, 2.28, 1.76, 0.87, color, True)
    text(slide, name, x+0.12, 2.54, 1.52, 0.24, 12.3, WHITE, True, PP_ALIGN.CENTER)
    if i < 5: text(slide, "→", x+1.77, 2.53, 0.28, 0.22, 16, MUTED, True, PP_ALIGN.CENTER)
rect(slide, 0.92, 4.06, 11.48, 1.47, PALE, True)
text(slide, "Trace attributes captured", 1.22, 4.36, 2.8, 0.25, 17, NAVY, True)
text(slide, "run ID • user/request type • timestamps • retrieval stage • generation attempts • candidate validation failures • retry reason • provider error • partial completion • final result", 1.22, 4.76, 10.35, 0.48, 15, INK)
text(slide, "This gives administrators a precise answer to: where did the request fail, and why?", 0.9, 6.2, 11.5, 0.3, 17, GREEN, True, PP_ALIGN.CENTER)

# 9 completion
slide = base("Report Completion Status — 50%", "Core pipeline is working; optimization and evaluation capabilities are the next development focus.", 9)
rect(slide, 0.68, 1.48, 5.81, 4.98, PALE, True)
rect(slide, 6.85, 1.48, 5.81, 4.98, WHITE, True, RGBColor(207,221,233))
text(slide, "Completed", 1.04, 1.86, 2.4, 0.31, 21, GREEN, True)
bullets(slide, ["Ingestion, cleaning, chunking, embedding, and retrieval pipeline", "LLM-based MCQ generation with diversity controls", "Structural, quality, source, and duplicate validation", "Partial results and explicit model/provider failure handling", "Monitoring page, trace storage, and API/model health indicators"], 1.02, 2.35, 4.95, 3.4, 15.5)
text(slide, "In Progress / Next", 7.21, 1.86, 3.4, 0.31, 21, ORANGE, True)
bullets(slide, ["HNSW and chunking parameter tuning", "Re-ranking and domain-specific embedding evaluation", "Second provider/model configuration for solver and critic", "RAG evaluation dashboard: relevance, faithfulness, precision, recall", "Semantic caching, memory policy, load and latency testing"], 7.2, 2.35, 4.95, 3.4, 15.5)

# 10 expected outcomes
slide = base("Expected Outcomes", "Target impact after the remaining optimization and evaluation work is completed.", 10)
outcomes = [
    ("Unlimited variation", "Generate many distinct questions from the same material using concept recombination and pattern rotation."),
    ("Higher assessment quality", "Grounded, meaningful 4-option MCQs with better distractors and fewer generic templates."),
    ("Measurable RAG quality", "Expose context relevance, faithfulness, retrieval precision/recall, duplicate rate, latency, and error rate."),
    ("Operational transparency", "Provide users and administrators exact, actionable reasons for unavailable models and failed requests.")]
for i, (h, d) in enumerate(outcomes):
    x = 0.68 + (i%2)*6.06; y = 1.64 + (i//2)*2.15
    rect(slide, x, y, 5.57, 1.62, PALE if i%2==0 else WHITE, True, RGBColor(207,221,233))
    text(slide, h, x+0.34, y+0.32, 4.8, 0.27, 18, NAVY, True)
    text(slide, d, x+0.34, y+0.79, 4.74, 0.52, 14.2, INK)

# 11 closing
slide = prs.slides.add_slide(prs.slide_layouts[6])
rect(slide, 0, 0, 13.333, 7.5, NAVY)
rect(slide, 0, 0, 13.333, 0.18, CYAN)
text(slide, "Conclusion", 0.86, 1.16, 7.5, 0.58, 31, WHITE, True)
text(slide, "The system combines RAG-based LLM generation with deterministic validation and traceable quality assurance.", 0.9, 2.06, 10.9, 0.68, 23, RGBColor(218, 232, 244))
rect(slide, 0.9, 3.34, 11.5, 1.43, BLUE, True)
text(slide, "Retrieval grounds the question in source evidence.\nGeneration creates diverse assessment candidates.\nValidation and monitoring make results dependable and explainable.", 1.25, 3.67, 10.8, 0.86, 19, WHITE, True, PP_ALIGN.CENTER)
text(slide, "Thank you", 0.9, 6.45, 3.5, 0.3, 18, CYAN, True)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT.resolve())
